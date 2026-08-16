"""
Sentinel AI — Book Pipeline Tests
==================================
Type: Unit tests for the pure-logic services behind the Manuscript (author)
and Publisher (manuscript) agents.

These cover the parts of the book pipeline that have no GUI and no LLM
dependency, so they run fast and deterministically:

  * services/book_exporter.py   — chapter detection, offsets, EPUB/DOCX/PDF export
  * services/content_calendar.py — posting-schedule generation
  * services/kdp_csv_parser.py  — KDP sales CSV summarisation
  * services/llm_parsing.py     — tolerant parsing of LLM list responses

Deliberately NOT covered here: anything that renders pixels (quote_graphics),
shells out to ffmpeg (shorts_generator), or calls a live LLM.

Run with:  pytest tests/test_book_pipeline.py -v
"""

import csv
import sys
import os
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from datetime import date
from pathlib import Path

from services.book_exporter import (
    split_into_chapters, find_chapter_offsets, export_book,
)
from services.content_calendar import build_calendar, CADENCE
from services.kdp_csv_parser import summarise_kdp_rows, parse_kdp_csv
from services.llm_parsing import parse_string_list


# ─────────────────────────────────────────────────────────────────────────────
# 1. Chapter detection
# Scenario: a draft is split into chapters for export and the Chapters navigator
# ─────────────────────────────────────────────────────────────────────────────

class TestSplitIntoChapters:

    def test_detects_numbered_chapters(self):
        text = "Chapter 1\nOpening line.\n\nChapter 2\nSecond line."
        chapters = split_into_chapters(text)
        assert len(chapters) == 2
        assert chapters[0][0] == "Chapter 1"
        assert chapters[1][0] == "Chapter 2"

    def test_body_is_attached_to_its_heading(self):
        text = "Chapter 1\nOpening line.\n\nChapter 2\nSecond line."
        chapters = split_into_chapters(text)
        assert "Opening line." in chapters[0][1]
        assert "Opening line." not in chapters[1][1]

    def test_detects_part_prologue_epilogue(self):
        text = "Prologue\nBefore.\n\nPart I\nMiddle.\n\nEpilogue\nAfter."
        headings = [h for h, _ in split_into_chapters(text)]
        assert headings == ["Prologue", "Part I", "Epilogue"]

    def test_detects_markdown_prefixed_headings(self):
        text = "## Chapter 1\nBody one.\n\n### Chapter 2\nBody two."
        headings = [h for h, _ in split_into_chapters(text)]
        assert headings == ["Chapter 1", "Chapter 2"]

    def test_is_case_insensitive(self):
        text = "CHAPTER 1: THE COLLAPSE\nBody.\n\nchapter 2: after\nMore."
        assert len(split_into_chapters(text)) == 2

    def test_text_before_first_heading_becomes_untitled_chapter(self):
        text = "Front matter here.\n\nChapter 1\nReal start."
        chapters = split_into_chapters(text)
        assert chapters[0][0] == ""
        assert "Front matter here." in chapters[0][1]
        assert chapters[1][0] == "Chapter 1"

    def test_draft_with_no_headings_is_one_chapter(self):
        # Export must still work on an unstructured draft — headings only improve it.
        text = "Just some prose with no chapter markers at all."
        chapters = split_into_chapters(text)
        assert len(chapters) == 1
        assert chapters[0][0] == ""
        assert chapters[0][1] == text

    def test_empty_text_yields_single_empty_chapter(self):
        assert split_into_chapters("") == [("", "")]

    def test_chapter_word_is_not_matched_mid_sentence(self):
        # "chapter" inside a sentence must not split the draft.
        text = "She opened the next chapter of her life and kept reading."
        assert len(split_into_chapters(text)) == 1


# ─────────────────────────────────────────────────────────────────────────────
# 2. Chapter offsets
# Scenario: the Chapters tab jumps the editor cursor to a chapter heading
# ─────────────────────────────────────────────────────────────────────────────

class TestFindChapterOffsets:

    def test_offset_points_at_the_heading_text(self):
        text = "Chapter 1\nBody one.\nChapter 2\nBody two."
        offsets = find_chapter_offsets(text)
        for off in offsets:
            assert text[off:].startswith("Chapter")

    def test_first_offset_is_zero_when_text_starts_with_heading(self):
        text = "Chapter 1\nBody."
        assert find_chapter_offsets(text) == [0]

    def test_offsets_are_in_document_order(self):
        text = "Chapter 1\nA.\nChapter 2\nB.\nChapter 3\nC."
        offsets = find_chapter_offsets(text)
        assert offsets == sorted(offsets)
        assert len(offsets) == 3

    def test_no_headings_yields_no_offsets(self):
        assert find_chapter_offsets("no headings here") == []

    def test_offset_count_matches_headed_chapters(self):
        # Pairing contract relied on by _author_refresh_chapters(): every chapter
        # with a heading consumes exactly one offset, in order.
        text = "Front matter.\n\nChapter 1\nA.\n\nChapter 2\nB."
        chapters = split_into_chapters(text)
        offsets = find_chapter_offsets(text)
        headed = [c for c in chapters if c[0]]
        assert len(offsets) == len(headed)


# ─────────────────────────────────────────────────────────────────────────────
# 3. Book export
# Scenario: a finished draft is exported to each submittable format
# ─────────────────────────────────────────────────────────────────────────────

class TestExportBook:

    SAMPLE = "Chapter 1\nThe first chapter body.\n\nChapter 2\nThe second chapter body."

    @pytest.mark.parametrize("fmt,ext", [("epub", "epub"), ("docx", "docx"), ("pdf", "pdf")])
    def test_export_creates_a_nonempty_file(self, tmp_path, fmt, ext):
        out = tmp_path / f"book.{ext}"
        result = export_book(self.SAMPLE, "My Book", "An Author", fmt, out)
        assert result == out
        assert out.exists()
        assert out.stat().st_size > 0

    def test_epub_is_a_valid_zip_with_expected_entries(self, tmp_path):
        import zipfile
        out = tmp_path / "book.epub"
        export_book(self.SAMPLE, "My Book", "An Author", "epub", out)
        with zipfile.ZipFile(out) as zf:
            names = zf.namelist()
        assert "mimetype" in names
        assert any(n.endswith("content.opf") for n in names)
        # One xhtml document per chapter
        assert sum(1 for n in names if n.endswith(".xhtml") and "chap_" in n) == 2

    def test_docx_contains_title_author_and_headings(self, tmp_path):
        from docx import Document
        out = tmp_path / "book.docx"
        export_book(self.SAMPLE, "My Book", "An Author", "docx", out)
        paragraphs = [p.text for p in Document(str(out)).paragraphs]
        assert "My Book" in paragraphs
        assert "An Author" in paragraphs
        assert "Chapter 1" in paragraphs

    def test_pdf_has_pages_and_extractable_text(self, tmp_path):
        import pypdf
        out = tmp_path / "book.pdf"
        export_book(self.SAMPLE, "My Book", "An Author", "pdf", out)
        reader = pypdf.PdfReader(str(out))
        assert len(reader.pages) >= 2  # title page + at least one chapter page
        assert "My Book" in reader.pages[0].extract_text()

    def test_unsupported_format_raises(self, tmp_path):
        with pytest.raises(ValueError):
            export_book(self.SAMPLE, "T", "A", "mobi", tmp_path / "book.mobi")

    def test_export_creates_missing_parent_directories(self, tmp_path):
        out = tmp_path / "nested" / "deeper" / "book.epub"
        export_book(self.SAMPLE, "My Book", "An Author", "epub", out)
        assert out.exists()

    def test_untitled_draft_still_exports(self, tmp_path):
        # Empty title/author fall back to placeholders rather than failing.
        out = tmp_path / "book.epub"
        export_book("Some prose.", "", "", "epub", out)
        assert out.exists()


# ─────────────────────────────────────────────────────────────────────────────
# 4. Content calendar
# Scenario: quotes are distributed into a posting schedule
# ─────────────────────────────────────────────────────────────────────────────

class TestBuildCalendar:

    QUOTES = [f"Quote {i}" for i in range(1, 6)]
    START = date(2026, 8, 3)

    def test_returns_empty_for_no_quotes(self):
        assert build_calendar([], 1, self.START, ["TikTok"]) == []

    def test_returns_empty_for_no_platforms(self):
        assert build_calendar(self.QUOTES, 1, self.START, []) == []

    def test_returns_empty_for_zero_weeks(self):
        assert build_calendar(self.QUOTES, 0, self.START, ["TikTok"]) == []

    def test_single_platform_matches_its_weekly_cadence(self):
        per_week, _ = CADENCE["TikTok"]
        slots = build_calendar(self.QUOTES, 1, self.START, ["TikTok"])
        assert len(slots) == per_week

    def test_two_weeks_doubles_the_post_count(self):
        one = build_calendar(self.QUOTES, 1, self.START, ["TikTok"])
        two = build_calendar(self.QUOTES, 2, self.START, ["TikTok"])
        assert len(two) == 2 * len(one)

    def test_slots_are_sorted_by_date(self):
        slots = build_calendar(self.QUOTES, 2, self.START, ["TikTok", "Instagram", "Pinterest"])
        assert [s.day for s in slots] == sorted(s.day for s in slots)

    def test_no_slot_falls_before_the_start_date(self):
        slots = build_calendar(self.QUOTES, 2, self.START, ["TikTok", "Instagram"])
        assert all(s.day >= self.START for s in slots)

    def test_no_slot_falls_outside_the_requested_window(self):
        weeks = 2
        slots = build_calendar(self.QUOTES, weeks, self.START, ["TikTok", "Instagram", "Pinterest"])
        last_allowed = self.START.toordinal() + weeks * 7 - 1
        assert all(s.day.toordinal() <= last_allowed for s in slots)

    def test_quotes_cycle_when_slots_outnumber_quotes(self):
        # Two quotes across a full multi-platform fortnight must not run out.
        slots = build_calendar(["A", "B"], 2, self.START, ["TikTok", "Instagram", "Pinterest"])
        assert len(slots) > 2
        assert set(s.quote for s in slots) == {"A", "B"}

    def test_tiktok_slots_are_all_shorts(self):
        slots = build_calendar(self.QUOTES, 1, self.START, ["TikTok"])
        assert {s.format for s in slots} == {"short"}

    def test_pinterest_slots_are_all_graphics(self):
        slots = build_calendar(self.QUOTES, 1, self.START, ["Pinterest"])
        assert {s.format for s in slots} == {"graphic"}

    def test_instagram_alternates_graphic_and_short(self):
        slots = build_calendar(self.QUOTES, 2, self.START, ["Instagram"])
        formats = {s.format for s in slots}
        assert formats == {"graphic", "short"}

    def test_every_slot_has_a_known_format(self):
        slots = build_calendar(self.QUOTES, 2, self.START, ["TikTok", "Instagram", "Pinterest"])
        assert all(s.format in ("graphic", "short") for s in slots)

    def test_captions_start_empty(self):
        # Captions are filled in by a later LLM step, not by the scheduler.
        slots = build_calendar(self.QUOTES, 1, self.START, ["TikTok"])
        assert all(s.caption == "" for s in slots)

    def test_unknown_platform_still_schedules_with_a_default(self):
        slots = build_calendar(self.QUOTES, 1, self.START, ["Threads"])
        assert len(slots) > 0
        assert all(s.platform == "Threads" for s in slots)


# ─────────────────────────────────────────────────────────────────────────────
# 5. KDP sales report summarisation
# Scenario: a downloaded KDP CSV is aggregated into headline metrics
# ─────────────────────────────────────────────────────────────────────────────

class TestSummariseKdpRows:

    def test_totals_units_and_royalties(self):
        rows = [
            {"Units Sold": "3", "Royalty": "6.00", "Marketplace": "Amazon.com"},
            {"Units Sold": "2", "Royalty": "4.00", "Marketplace": "Amazon.com"},
        ]
        summary = summarise_kdp_rows(rows)
        assert summary["total_units"] == 5
        assert summary["total_royalties_usd"] == 10.00

    def test_groups_by_marketplace(self):
        rows = [
            {"Units Sold": "3", "Royalty": "6.00", "Marketplace": "Amazon.com"},
            {"Units Sold": "1", "Royalty": "2.00", "Marketplace": "Amazon.co.uk"},
        ]
        by_market = {m["marketplace"]: m for m in summarise_kdp_rows(rows)["by_marketplace"]}
        assert by_market["Amazon.com"]["units"] == 3
        assert by_market["Amazon.co.uk"]["units"] == 1

    def test_sums_kenp_pages_when_present(self):
        rows = [
            {"Units Sold": "0", "Royalty": "0", "Marketplace": "Amazon.com", "KENP Read": "1200"},
            {"Units Sold": "0", "Royalty": "0", "Marketplace": "Amazon.com", "KENP Read": "800"},
        ]
        assert summarise_kdp_rows(rows)["kenp_pages_read"] == 2000

    def test_missing_columns_default_to_zero(self):
        # KDP column names vary by report type — missing ones must not crash.
        rows = [{"Marketplace": "Amazon.com"}]
        summary = summarise_kdp_rows(rows)
        assert summary["total_units"] == 0
        assert summary["total_royalties_usd"] == 0.0

    def test_malformed_rows_are_skipped_not_fatal(self):
        rows = [
            {"Units Sold": "abc", "Royalty": "xyz", "Marketplace": "Amazon.com"},
            {"Units Sold": "2", "Royalty": "5.00", "Marketplace": "Amazon.com"},
        ]
        summary = summarise_kdp_rows(rows)
        assert summary["total_units"] == 2
        assert summary["total_royalties_usd"] == 5.00

    def test_empty_rows_yield_zeroed_summary(self):
        summary = summarise_kdp_rows([])
        assert summary["total_units"] == 0
        assert summary["by_marketplace"] == []

    def test_royalties_are_rounded_to_cents(self):
        rows = [{"Units Sold": "1", "Royalty": "3.333", "Marketplace": "Amazon.com"}]
        assert summarise_kdp_rows(rows)["total_royalties_usd"] == 3.33

    def test_parse_kdp_csv_reads_and_strips_fields(self, tmp_path):
        csv_path = tmp_path / "report.csv"
        with open(csv_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Units Sold", " Royalty ", "Marketplace"])
            writer.writerow([" 4 ", " 8.00 ", " Amazon.com "])
        rows = parse_kdp_csv(csv_path)
        assert rows[0]["Units Sold"] == "4"
        assert rows[0]["Royalty"] == "8.00"   # header whitespace stripped too


# ─────────────────────────────────────────────────────────────────────────────
# 6. LLM list-response parsing
# Scenario: quote/caption responses arrive in whatever shape the model chose
# ─────────────────────────────────────────────────────────────────────────────

class TestParseStringList:

    def test_parses_clean_json_array(self):
        assert parse_string_list('["one", "two"]') == ["one", "two"]

    def test_strips_markdown_json_fence(self):
        assert parse_string_list('```json\n["one", "two"]\n```') == ["one", "two"]

    def test_strips_bare_markdown_fence(self):
        assert parse_string_list('```\n["one", "two"]\n```') == ["one", "two"]

    def test_extracts_json_array_embedded_in_prose(self):
        raw = 'Sure! Here are the quotes:\n["one", "two"]\nHope that helps.'
        assert parse_string_list(raw) == ["one", "two"]

    def test_falls_back_to_numbered_lines(self):
        assert parse_string_list("1. First\n2. Second") == ["First", "Second"]

    def test_falls_back_to_bulleted_lines(self):
        assert parse_string_list("- First\n- Second") == ["First", "Second"]

    def test_fallback_strips_surrounding_quote_marks(self):
        assert parse_string_list('- "First"\n- "Second"') == ["First", "Second"]

    def test_drops_blank_entries(self):
        assert parse_string_list('["one", "", "  ", "two"]') == ["one", "two"]

    def test_empty_response_yields_empty_list(self):
        assert parse_string_list("") == []

    def test_always_returns_a_list_never_raises(self):
        # The caller renders this straight into a list widget; it must never throw.
        for junk in ["{not: valid}", "[[[", "null", "42", "   "]:
            assert isinstance(parse_string_list(junk), list)


# ─────────────────────────────────────────────────────────────────────────────
# 7. Asset output paths
# Scenario: several graphics/shorts are generated in quick succession
# ─────────────────────────────────────────────────────────────────────────────

class TestUniqueOutputPath:
    """Regression guard: asset filenames used to be timestamp-only, so generating
    two graphics inside the same tick silently overwrote the first."""

    def test_rapid_calls_never_collide(self, tmp_path):
        from ui.book_widgets import unique_output_path
        paths = []
        for _ in range(10):
            p = unique_output_path(tmp_path, "quote", ".png")
            p.write_bytes(b"x")          # occupy it, as a real render would
            paths.append(p)
        assert len(set(paths)) == 10

    def test_respects_stem_and_suffix(self, tmp_path):
        from ui.book_widgets import unique_output_path
        p = unique_output_path(tmp_path, "short", ".mp4")
        assert p.name.startswith("short_")
        assert p.suffix == ".mp4"

    def test_creates_missing_directory(self, tmp_path):
        from ui.book_widgets import unique_output_path
        target = tmp_path / "nested" / "dir"
        p = unique_output_path(target, "quote", ".png")
        assert target.exists()
        assert p.parent == target

    def test_paired_png_and_mp4_share_a_stem(self, tmp_path):
        # Shorts derive the image path from the video path, so they must match.
        from ui.book_widgets import unique_output_path
        video = unique_output_path(tmp_path, "short", ".mp4")
        image = video.with_suffix(".png")
        assert video.stem == image.stem
